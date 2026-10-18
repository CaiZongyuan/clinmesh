"""Add independently sourced prices and use-specific status labels to a new offline packet."""

import argparse
import collections
import json
import re
import shutil
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
import pymupdf
import xlrd

from verify import digest, load, rows, verify


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_rows(path, values):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as out:
        for value in values:
            out.write(json.dumps(value, ensure_ascii=False, separators=(",", ":")) + "\n")


def compact(value):
    return re.sub(r"\s+", "", str(value or ""))


def identifier(value):
    if value is None or value == "":
        return None
    if isinstance(value, str):
        return value
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer() and abs(value) <= 2**53 - 1:
        return str(int(value))
    raise ValueError(f"Identifier cannot be recovered faithfully: {value!r}")


def money(value):
    if isinstance(value, bool) or value is None:
        raise ValueError("Missing or nonnumeric price")
    raw = str(value).strip()
    if not re.fullmatch(r"\d+(?:\.\d+)?", raw):
        raise ValueError(f"Unrecognised monetary literal: {raw!r}")
    try:
        result = Decimal(raw)
    except InvalidOperation as error:
        raise ValueError(raw) from error
    if not result.is_finite() or result < 0:
        raise ValueError(raw)
    return format(result.normalize(), "f")


def explicit_package(raw):
    """Read a single explicit count/unit/unit expression; do not invent nested conversions."""
    match = re.fullmatch(r"(\d+)\*?([^\d\s*/×]+)[/]([^\d\s*/×]+)", str(raw or "").strip())
    if not match or int(match[1]) < 1:
        return None
    return {"count": int(match[1]), "contentUnit": match[2], "packageUnit": match[3],
            "basis": "explicit_source_package_text", "hospitalConfirmed": False}


def excel(path):
    with path.open("rb") as source:
        book = openpyxl.load_workbook(source, read_only=True, data_only=False)
        sheet = book.worksheets[0]
        title = sheet.title
        values = list(sheet.values)
        book.close()
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(path) as source:
        sheet_xml = ET.fromstring(source.read("xl/worksheets/sheet1.xml"))
    lexical = {c.attrib["r"]: c.find("s:v", ns).text for c in sheet_xml.findall(".//s:c", ns)
               if c.find("s:v", ns) is not None and c.attrib.get("t", "n") == "n"}
    return title, values, lexical


def source_price_status(record):
    if record.get("currentApplicability") == "superseded_historical":
        return {"label": "不可作当前价", "reason": "上一轮浙江冠脉支架报价已被第二轮接续替代；仅保留历史参考"}
    missing = []
    if not record["price"]["unit"]:
        missing.append("价格计量单位或金额分母")
    if not record.get("specification"):
        missing.append("具体规格型号")
    return {"label": "待核实", "reason": "地区、当前有效期和本院适用性尚未确认", "missingFields": missing}


def procedure_use(record):
    if record.get("nameConflict"):
        return "暂不可用（隔离）", "微波与射频术式冲突，不能自动选择正确值"
    reason = "动态条目的有效状态、生效日期及DRG配套版本缺少依据"
    if record.get("nameReviewRequired"):
        reason += "；DRG各处来源名称存在文字变体"
    return "待核实", reason


class Enrichment:
    def __init__(self, candidate, output, desktop, evidence):
        self.base, self.out = Path(candidate).resolve(), Path(output).resolve()
        self.desktop, self.evidence = Path(desktop).resolve(), Path(evidence).resolve()
        if self.out.exists() or self.out.is_relative_to(self.base):
            raise ValueError("Output must be a new directory outside the previous packet")
        verify(self.base)
        shutil.copytree(self.base, self.out)
        shutil.copyfile(self.base / "checksums.json", self.out / "review/round-2-checksums.json")
        self.manifest = load(self.out / "release-candidate.json")
        self.downloads = load(self.evidence / "source-download-results.json")
        self.meta = {r["id"]: r for r in load(self.evidence / "source-metadata.json")}
        self.sources = {}
        self.tables = {}

    def table(self, name, values):
        write_rows(self.out / "standardized" / (name + ".jsonl"), values)
        self.manifest["tables"][name] = len(values)
        self.tables[name] = values

    def preserve_sources(self):
        for entry in self.downloads:
            assert entry["status"] == "hash_matched", entry["id"]
            path = Path(entry["path"])
            assert digest(path) == entry["sha256"] == entry["expectedSha256"]
            dest = self.out / "raw/price-sources/official" / path.name
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(path, dest)
            self.sources[entry["id"]] = {k: v for k, v in entry.items() if k != "path"}
            self.sources[entry["id"]].update(path=dest.relative_to(self.out).as_posix(),
                                            metadata=self.meta[entry["id"]])
        for pattern, source_id in [("江西*.xls", "JX-LOCAL"), ("附件：广东*.xlsx", "GD-LOCAL"),
                                   ("心血管*.xlsx", "CARDIO-LOCAL")]:
            matches = list(self.desktop.glob(pattern))
            assert len(matches) == 1
            original = matches[0]
            dest = self.out / "raw/price-sources/provided" / original.name
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(original, dest)
            self.sources[source_id] = {"id": source_id, "path": dest.relative_to(self.out).as_posix(),
                                       "sha256": digest(dest), "bytes": dest.stat().st_size,
                                       "status": "provided_local_file"}
        for local_id, official_id in [("JX-LOCAL", "JX-DRUG-20250930"), ("GD-LOCAL", "GD-DRUG-20260930")]:
            local = excel(self.out / self.sources[local_id]["path"])[1]
            official = excel(self.out / self.sources[official_id]["path"])[1]
            assert local == official
            self.sources[local_id]["cellMatrixEqualToOfficial"] = official_id
            self.sources[local_id]["binaryEqualToOfficial"] = self.sources[local_id]["sha256"] == self.sources[official_id]["sha256"]
        write_json(self.out / "review/price-source-manifest.json", self.sources)

    def source(self, source_id, **locator):
        s = self.sources[source_id]
        return {"kind": "external_price_file", "sourceId": source_id, "path": s["path"],
                "sha256": s["sha256"], "url": s.get("url"), **locator}

    def drugs(self, source_id, region, header_row, expected):
        path = self.out / self.sources[source_id]["path"]
        sheet, data, lexical = excel(path)
        header = data[header_row - 1]
        is_jx = source_id.startswith("JX")
        price_column = "N" if is_jx else "M"
        records = []
        for row_number, raw in enumerate(data[header_row:], header_row + 1):
            if raw[0] is None:
                continue
            assert all(not (isinstance(v, str) and v.startswith("=")) for v in raw)
            assert int(raw[0]) == len(records) + 1
            fields = dict(zip(header, raw))
            literal = lexical.get(f"{price_column}{row_number}") or str(raw[13 if is_jx else 12])
            amount = money(literal)
            packaging = raw[9] if is_jx else raw[7]
            package = explicit_package(packaging)
            code = identifier(raw[2 if is_jx else 3])
            record = {"id": f"price-{source_id}-{row_number:04}", "productCode": code,
                      "codeSystem": "source_drug_unified_code" if is_jx else "source_nhsa_drug_code",
                      "platformProductId": None if is_jx else identifier(raw[2]),
                      "name": raw[3 if is_jx else 4], "dosageForm": raw[4 if is_jx else 5],
                      "specification": raw[5 if is_jx else 6], "manufacturer": raw[11 if is_jx else 9],
                      "declaringEnterprise": raw[12 if is_jx else 10],
                      "approvalNumber": raw[14 if is_jx else 11], "packageRaw": packaging,
                      "packageStructure": package, "conversionRatioRaw": raw[6] if is_jx else None,
                      "smallestPreparationUnit": raw[7] if is_jx else None,
                      "smallestPackageUnit": raw[8] if is_jx else None,
                      "price": {"kind": "selected_procurement_price" if is_jx else "listed_package_price",
                                "amount": amount, "currency": "CNY", "raw": literal,
                                "basis": "denominator_unverified" if is_jx else "per_source_package",
                                "unit": None if is_jx else package["packageUnit"] if package else None},
                      "region": region, "snapshotDate": "2025-09-30" if is_jx else None,
                      "batch": raw[1], "publicationDate": "2025-10-14" if is_jx else "2026-10-13",
                      "validFrom": None, "validTo": None, "currentApplicability": "unverified",
                      "rawFields": fields, "source": self.source(source_id, worksheet=sheet,
                                                                            row=row_number, priceCell=f"{price_column}{row_number}"),
                      "hospitalMapping": {"state": "待补全", "hospitalProductId": None,
                                          "chargePrice": None, "procurementCost": None, "unitConversion": None},
                      "runtimeEnabled": False}
            record["currentPriceUse"] = source_price_status(record)
            if is_jx:
                record["currentPriceUse"]["reason"] += "；原价格列仅标中选价(元)，未据包装列推定金额分母"
            records.append(record)
        assert len(records) == expected
        self.table("prices/jiangxi-drugs" if is_jx else "prices/guangdong-drugs", records)

    def devices(self):
        sheet, values, _ = excel(self.out / self.sources["CARDIO-LOCAL"]["path"])
        source_book = xlrd.open_workbook(self.out / self.sources["HN-PACING-2026"]["path"])
        pacing = source_book.sheet_by_index(0)
        pdf_rows = {}
        for source_id in self.meta:
            if source_id == "HN-PACING-2026" or source_id.endswith("DRUG-20250930") or source_id.endswith("DRUG-20260930"):
                continue
            with pymupdf.open(self.out / self.sources[source_id]["path"]) as pdf:
                pdf_rows[source_id] = [(page_no, ti, ri, raw) for page_no, page in enumerate(pdf, 1)
                                       for ti, t in enumerate(page.find_tables().tables)
                                       for ri, raw in enumerate(t.extract()) if ri > 0]
        records, corrections = [], []
        for row_number, r in enumerate(values[7:], 8):
            if r[0] is None:
                continue
            assert all(not (isinstance(v, str) and v.startswith("=")) for v in r)
            source_id = r[13]
            amount = money(r[4])
            registrant = None
            enterprise, code_raw = r[2], identifier(r[11])
            code = code_raw
            flags = []
            code_candidates = []
            if source_id == "HN-PACING-2026":
                source_row = int(re.fullmatch(r"Sheet1 第(\d+)行", r[14])[1])
                raw = pacing.row_values(source_row - 1)
                assert compact(raw[2]) == compact(r[1]) and compact(raw[3]) == compact(r[2])
                assert money(raw[8]) == amount and raw[7] == r[5]
                assert (raw[5] if raw[5] != "/" else None) == r[11]
                assert (money(raw[9]) if raw[9] != "/" else None) == (money(r[16]) if r[16] is not None else None)
                registrant = raw[4]
                locator = {"worksheet": pacing.name, "row": source_row, "priceCell": f"I{source_row}"}
            else:
                candidates = []
                for page, ti, ri, raw in pdf_rows[source_id]:
                    text = compact(" ".join(str(v or "") for v in raw))
                    if compact(r[12]) in text and money(raw[-1]) == amount and compact(r[1]) in text:
                        candidates.append((page, ti, ri, raw))
                assert len(candidates) == 1, (source_id, r[1], len(candidates))
                page, ti, ri, raw = candidates[0]
                locator = {"pdfPage": page, "table": ti, "row": ri}
                if code_raw and "\n" in code_raw:
                    code_candidates = code_raw.splitlines()
                    code = None
                    flags.append("multiple_codes_and_source_cell_overflow")
                    corrections.append({"recordId": f"price-CARDIO-{row_number:04}", "field": "productCode",
                                        "originalValue": code_raw, "correctedValue": None,
                                        "candidateValues": code_candidates, "reviewStatus": "pending_source_layout",
                                        "reason": "同一报价含多码且第三码末尾跨入邻行；保留全部候选，不选定一个产品码",
                                        "source": self.source(source_id, **locator), "rawSourceRow": raw})
                elif code_raw:
                    assert compact(code_raw) in compact(" ".join(str(v or "") for v in raw))
                source_enterprise = compact(raw[-2]) if source_id.startswith("TJ") else compact(raw[3])
                if compact(enterprise) != source_enterprise:
                    complete = source_enterprise.endswith("公司")
                    corrections.append({"recordId": f"price-CARDIO-{row_number:04}", "field": "declaringEnterprise",
                                        "originalValue": enterprise, "correctedValue": source_enterprise if complete else None,
                                        "candidateValues": [enterprise, source_enterprise],
                                        "reviewStatus": "resolved_official_pdf_cell" if complete else "pending_source_layout",
                                        "reason": "对照原页移除错接的前行字形" if complete else "原PDF单元格企业名仍截断；不猜补末字",
                                        "source": self.source(source_id, **locator), "rawSourceRow": raw})
                    enterprise = source_enterprise if complete else None
                    if not complete:
                        flags.append("enterprise_name_truncated_in_source_layout")
                if source_id.startswith("ZJ"):
                    registrant = raw[1] if source_id.endswith("1") else raw[4]
            unit = None if r[5] == "原表未列" else r[5]
            record = {"id": f"price-CARDIO-{row_number:04}", "name": r[1], "category": r[0],
                      "declaringEnterprise": enterprise, "providedEnterpriseRaw": r[2], "registrant": registrant, "specification": r[3],
                      "productCode": code, "codeSystem": "nhsa_consumable_27_character_code" if code and len(code) == 27 else
                      "source_consumable_20_character_code" if code else None,
                      "providedCodeRaw": code_raw, "candidateProductCodes": code_candidates, "reviewFlags": flags,
                      "registrationNumber": r[12], "packageRaw": None,
                      "price": {"kind": "selected_procurement_price", "sourceKindRaw": r[6],
                                "amount": amount, "currency": "CNY", "unit": unit, "raw": str(r[4])},
                      "proposedSupplyPrice": money(r[16]) if r[16] is not None else None,
                      "region": r[7], "validFrom": str(r[8])[:10], "validTo": str(r[9])[:10] if r[9] else None,
                      "currentApplicability": "superseded_historical" if source_id.startswith("ZJ") else "unverified",
                      "applicabilityRaw": r[10], "notes": r[15], "originalSourceRow": raw,
                      "source": self.source(source_id, **locator),
                      "providedWorkbook": self.source("CARDIO-LOCAL", worksheet=sheet, row=row_number),
                      "hospitalMapping": {"state": "待补全", "hospitalProductId": None, "chargePrice": None,
                                          "procurementCost": None, "unitConversion": None}, "runtimeEnabled": False}
            record["currentPriceUse"] = source_price_status(record)
            if flags:
                record["currentPriceUse"]["reason"] += "；企业名称或多码排版仍待核实，不能用于自动产品匹配"
            records.append(record)
        assert len(records) == 231
        self.table("prices/cardiovascular", records)
        write_rows(self.out / "review/price-source-corrections.jsonl", corrections)
        for name in ["price-source-layout-review.png", "price-source-row-crops.png"]:
            destination = self.out / "raw/price-sources/layout-evidence" / name
            destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(self.evidence / name, destination)

    def readiness(self):
        result = []
        primary = ["icd10/records", "icd9/records", "services/records", "drugs/records", "drg/groups"]
        for table in primary:
            for r in rows(self.out / "standardized" / (table + ".jsonl")):
                dataset = table.split("/")[0]
                heading = r.get("kind") == "hierarchy_only" or r.get("entryKind") == "catalogue_heading"
                status = {"id": f"readiness-{r['id']}", "table": table, "recordId": r["id"],
                          "code": r.get("code") or r.get("catalogueNumber"), "name": r.get("name"),
                          "sourceReference": "仅作分类" if heading else "可作来源参考",
                          "currentApplicability": "待核实", "runtimeEnabled": False}
                if dataset in ["services", "drugs"]:
                    status["hospitalUse"] = "不适用" if heading else "待补全"
                    status["hospitalUseReason"] = "分类标题不能作为收费末级" if heading else "未取得本院项目/产品主数据、收费价及成本有效版本"
                    status["missingFields"] = [] if heading else ["院内项目或产品编码", "院内计价单位及换算", "本院收费价", "采购或归集成本", "有效起止日", "确认依据"]
                    if dataset == "services":
                        status["sourcePriceUse"] = "不适用" if heading else "待补全" if r["reviewStatus"] == "waiting_external_pricing" else "可作2024来源参考"
                        if not heading and not r.get("billingUnit"):
                            status["missingFields"].append("来源计价单位")
                    else:
                        p = r["paymentStandard"]
                        status["paymentStandardUse"] = "保密不可取得" if p["state"] == "confidential" else "待补全" if p["state"] == "unknown" else "按原规格条件参考"
                        status["paymentStandardState"] = p["state"]
                else:
                    status["clinicalGroupingUse"] = "不可作为末级编码" if heading else "待核实"
                    status["reason"] = "匹配版本、地方适用性及规则语义未完成审核；来源文字核对不等于医学审定"
                result.append(status)
        procedure = load(self.out / "review/procedure-code-review.json")
        for r in procedure["records"]:
            use, reason = procedure_use(r)
            result.append({"id": f"readiness-procedure-{r['code']}", "table": "review/procedure-code-review",
                           "recordId": r["code"], "code": r["code"], "sourceReviewState": r["state"],
                           "sourceReference": "可作来源参考", "clinicalGroupingUse": use, "reason": reason,
                           "runtimeEnabled": False})
        for table, records in self.tables.items():
            if not table.startswith("prices/"):
                continue
            for r in records:
                result.append({"id": "readiness-" + r["id"], "table": table, "recordId": r["id"],
                               "code": r["productCode"], "name": r["name"],
                               "sourceReference": "待核实" if r.get("reviewFlags") else "可作来源价格参考",
                               "sourceReviewFlags": r.get("reviewFlags", []),
                               "currentPriceUse": r["currentPriceUse"], "hospitalUse": "待补全",
                               "reason": "无院内目录和价格，不建立已确认的院内映射", "runtimeEnabled": False})
        self.table("assessment/record-readiness", result)
        write_json(self.out / "review/status-definitions.json", {
            "可作来源参考": "来源事实可检索、追溯；不表示当前院内适用",
            "待补全": "需要的事实未提供，字段保持空值；补全必须带来源及有效期",
            "待核实": "已有候选值，但版本、单位、名称或适用性缺少充分证据",
            "暂不可用（隔离）": "存在明确冲突，保留原文，不能参与对应自动用途",
            "不可作当前价": "已确认仅为被接续替代的历史价格；历史来源仍保留",
            "不适用": "分类标题不能作为可收费末级；仍可用来组织目录",
            "保密不可取得": "来源以*表示保密，不猜测、不填零；区别于普通缺失",
            "scope": "标签按用途分别判断；同一条可作来源参考，同时院内用途待补全"})

    def reports(self):
        counts = {key: len(value) for key, value in self.tables.items()}
        prices = [r for k, values in self.tables.items() if k.startswith("prices/") for r in values]
        duplicates = []
        for table, values in self.tables.items():
            if not table.startswith("prices/"):
                continue
            by_code = collections.defaultdict(list)
            for r in values:
                if r["productCode"]:
                    by_code[r["productCode"]].append(r)
            for code, matching in by_code.items():
                if len(matching) > 1:
                    duplicates.append({"table": table, "code": code, "recordIds": [r["id"] for r in matching],
                                       "label": "待核实", "reason": "原码重复保留，不能用编码字典覆盖行；需联合包装、企业、注册证及价格批次"})
        write_rows(self.out / "review/price-duplicate-codes.jsonl", duplicates)
        price_corrections = list(rows(self.out / "review/price-source-corrections.jsonl"))
        summary = {"round": 3, "scope": "offline_data_only", "runtimeEnabled": False, "counts": counts,
                   "officialPriceSourcesHashVerified": len(self.downloads), "priceRecords": len(prices),
                   "sourcePricesMissingUnit": sum(not r["price"]["unit"] for r in prices),
                   "supersededCurrentPrices": sum(r["currentApplicability"] == "superseded_historical" for r in prices),
                   "confirmedHospitalMappings": 0, "duplicateCodeGroupsPreserved": len(duplicates),
                   "priceTextReviewOutcomes": dict(collections.Counter(r["reviewStatus"] for r in price_corrections)),
                   "procedurePending": {"nameVariants": 17, "clinicalConflictQuarantined": 1, "officialDetailsMissing": 27},
                   "servicesWaitingExternalPricing": 72, "servicesWaitingUnit": 22,
                   "priceSourcePolicy": "amounts stay in separate source/region/batch price records; never fill hospital prices",
                   "previousChecksumsSha256": digest(self.base / "checksums.json")}
        write_json(self.out / "reports/round-3-data-status.json", summary)
        guide = Path(__file__).with_name("hospital-mapping-guide.md")
        shutil.copyfile(guide, self.out / "reports/hospital-mapping-guide.md")
        report = f"""# 数据补充与使用状态（r3）

本轮只处理数据。五包原文、字段修正、条件价格及r2复核证据继续保留。院内价格、有效期和编码适用性没有充分依据的值保持为空或待核实。

## 已补充的参考价格

| 来源 | 记录数 | 金额类型 | 使用状态 |
| --- | ---: | --- | --- |
| 江西2025-09-30快照 | 1974 | 集采中选价 | 可作来源参考；价格分母、当前执行及院内适用性待核实 |
| 广东20260930批次 | 78 | 包装挂网价 | 可作指定批次参考；当前执行及院内适用性待核实 |
| 湖南2026起搏、除颤及导线 | 52 | 集采中选价 | 已补原附件中的医疗器械注册人；规格及院内映射待补全 |
| 天津2023冠脉四类 | 163 | 历史公布中选价 | 原表缺逐行单位及规格；接续、调价和当前适用性待核实 |
| 浙江上一轮冠脉支架 | 16 | 上一轮中选价 | 不可作当前价，保留历史参考 |

共{len(prices)}条，分别保存在`standardized/prices/`。9份官方附件均重新取得并通过既有来源哈希核对。桌面江西表与官方原件字节不同，但逐单元格矩阵完全一致；两份均保留，标准化以官方原件为依据。金额不并入本院收费或成本。

湖南52条原附件额外补入医疗器械注册人；其中35条有27位基准品（含替代）耗材码，17条缺码。天津原报价以20位代码列示，不能补零变成27位规格产品码；1条含多码及跨行排版，保留全部候选但不选定单个产品码。所有231条耗材报价仍缺具体规格型号，179条还缺逐行计价单位。

原页复核另纠正5处企业名错接字符；5处企业名在原PDF单元格仍截断、1处多码布局进入待核实。它们保存于`review/price-source-corrections.jsonl`，原值、候选值、原表行和页码分别保留；待核实的企业标准值或单一产品码留空，不猜补。

江西价格列仅称“中选价(元)”，本轮没有把最小包装单位推定为金额分母，也未据此计算每片/每瓶价格。显式包装结构单独保存；复杂包装不强行拆分。所有金额为十进制字符串，原字段及价格单元格定位保留。重复编码共{len(duplicates)}组，逐行保留并列入`review/price-duplicate-codes.jsonl`，不自动合并。

## 待补全、待核实与暂不可用

| 对象或用途 | 当前状态 | 后续需要什么 |
| --- | --- | --- |
| 72条服务来源缺价 | 待补全 | 47条院内自主定价，25条适用地区或院内有效价；22条另缺单位 |
| 全部院内项目、药品与耗材映射 | 待补全 | 本院主数据、收费价、采购/归集成本、包装换算及有效期；已确认映射数为0 |
| 17个操作码名称变体 | 待核实 | 权威名称或勘误依据；原文与候选均保留 |
| 1个微波/射频术式冲突 | 暂不可用（隔离） | 权威勘误或配套字典确认 |
| 27个操作码官方详情未取得 | 待核实 | 可验收官方详情或匹配版本字典；未把请求失败认作缺码 |
| 53个在线操作码的有效状态及DRG配套版本 | 待核实 | 发布批次、生效日期、有效状态与兼容声明 |
| 分类标题、纯层级行 | 仅作分类，不可作收费/编码末级 | 保留子项及原层级，不制造末级码或价格 |
| 药品保密支付标准* | 保密不可取得 | 保持*，不推算金额，不改成0 |

状态按用途保存于`standardized/assessment/record-readiness.jsonl`，含原表名、稳定记录ID、具体原因及缺失字段；定义见`review/status-definitions.json`。一个项目可供来源查阅，同时院内使用仍待补全。这里的“可作参考”不代表当前适用或已经人工医学审定。

72条服务逐项位置见[服务复核速览](review-overview.md)，80个操作码逐项证据见[跨库报告](cross-library-report.md)。[院内映射说明](hospital-mapping-guide.md)解释需要的表和匹配方法。

本轮未新增院内映射，不把医保药品目录编号当产品码，不将ICD操作码直接当收费项目码；没有接入参考数据库或动态DRG流程。r2原始标准表内容不变，使用状态以独立关联表补充。检查结果另存交付目录旁的verification.json。
"""
        (self.out / "reports/data-status.md").write_text(report, encoding="utf-8")
        (self.out / "README.md").write_text("# ClinMesh 参考数据候选包 r3\n\n先看[数据补充与使用状态](reports/data-status.md)和[院内映射说明](reports/hospital-mapping-guide.md)。\n\n五类原数据、r2复核和原始档案继续保留；新增参考价格与按用途判断的状态。标准数据为JSONL，编码是字符串。未接入运行库。\n", encoding="utf-8")

    def finish(self):
        old_release = self.manifest["releaseId"]
        self.manifest.update(releaseId="clinmesh-five-catalogues-2026-10-16.r3",
                             previousCandidate={"releaseId": old_release, "checksumsSha256": digest(self.base / "checksums.json")},
                             dataStatusReport="reports/data-status.md", priceSourceManifest="review/price-source-manifest.json")
        write_json(self.out / "release-candidate.json", self.manifest)
        for name in ["enrich.py", "test_enrich.py", "verify.py", "hospital-mapping-guide.md", "requirements.txt"]:
            shutil.copyfile(Path(__file__).with_name(name), self.out / "tooling" / name)
        files = [{"path": p.relative_to(self.out).as_posix(), "bytes": p.stat().st_size, "sha256": digest(p)}
                 for p in sorted(self.out.rglob("*")) if p.is_file() and p != self.out / "checksums.json"]
        write_json(self.out / "checksums.json", files)

    def run(self):
        self.preserve_sources()
        self.drugs("JX-DRUG-20250930", "江西省", 2, 1974)
        self.drugs("GD-DRUG-20260930", "广东省", 3, 78)
        self.devices()
        self.readiness()
        self.reports()
        self.finish()
        print(json.dumps(load(self.out / "reports/round-3-data-status.json"), ensure_ascii=False), flush=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--candidate", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--desktop", type=Path, required=True)
    parser.add_argument("--evidence", type=Path, required=True)
    args = parser.parse_args()
    Enrichment(args.candidate, args.output, args.desktop, args.evidence).run()
